"""샌프란시스코 연준 수집기 — 수요·공급 요인별 PCE 인플레이션 (Shapiro 분해).

Adam Shapiro(SF 연준)가 매달 갱신하는 지표다. PCE 바스켓의 각 품목에 대해
가격과 수량의 '예상 밖 변화'가 같은 방향이면 수요 주도, 반대 방향이면 공급 주도,
통계적으로 구분이 안 되면 모호(ambiguous)로 분류한다.
  · 10년 롤링 회귀로 그달의 예측값을 만들고 실제값과 비교한다.
  · 세 갈래 기여도를 더하면 PCE 인플레이션이 된다.

받는 곳은 두 가지다.
  엑셀 (기본)  supply-demand-pce-inflation.xlsx — 'Data' 시트 한 장에 12열.
               1969년부터 있다. 이걸 쓴다.
  CSV (예비)   차트별 파일 4개. 최근 5년만 담고 있어 엑셀이 막혔을 때만 쓴다.

indicators.yaml 사용 예:
  - id: us_pce_sd
    name: 미국 PCE 수요·공급 분해
    source: frbsf
    unit: '%p'
    freq: M
    merge_always: true
    params:
      xlsx: supply-demand-pce-inflation.xlsx
      local: data/supply-demand-pce-inflation.xlsx   # 다운로드 실패 시 이 파일
      files:                                          # 엑셀이 안 될 때만
        - {csv: supply-demand-pce-core-yoy-chart-4.csv, prefix: '근원 YoY'}
"""
import io
import csv
import re
from datetime import date
from pathlib import Path

import requests

BASE = "https://www.frbsf.org/wp-content/uploads/"
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}
ROOT = Path(__file__).resolve().parent.parent

# 엑셀 열 이름 → 표시명. 원본 열 이름이 바뀌면 여기만 고치면 된다.
XL_COLS = {
    "Demand-driven Inflation (core, y/y)": "근원 YoY · 수요",
    "Supply-driven Inflation (core, y/y)": "근원 YoY · 공급",
    "Ambiguous (core, y/y)": "근원 YoY · 모호",
    "Demand-driven Inflation (headline, y/y)": "헤드라인 YoY · 수요",
    "Supply-driven Inflation (headline, y/y)": "헤드라인 YoY · 공급",
    "Ambiguous (headline, y/y)": "헤드라인 YoY · 모호",
    "Demand-driven Inflation (core, m/m)": "근원 월간연율 · 수요",
    "Supply-driven Inflation (core, m/m)": "근원 월간연율 · 공급",
    "Ambiguous (core, m/m)": "근원 월간연율 · 모호",
    "Demand-driven Inflation (headline, m/m)": "헤드라인 월간연율 · 수요",
    "Supply-driven Inflation (headline, m/m)": "헤드라인 월간연율 · 공급",
    "Ambiguous (headline, m/m)": "헤드라인 월간연율 · 모호",
}
# CSV 열 이름 → 표시명 (예비 경로)
CSV_COLS = {
    "Demand-driven Inflation": "수요",
    "Supply-driven Inflation": "공급",
    "Ambiguous": "모호",
}


class FrbsfError(RuntimeError):
    pass


def _to_date(s) -> str | None:
    """'2026m7' → '2026-07-31' (월말). 앞뒤 공백이 붙어 온다."""
    m = re.fullmatch(r"\s*(\d{4})m(\d{1,2})\s*", str(s or ""))
    if not m:
        return None
    y, mo = int(m.group(1)), int(m.group(2))
    if not 1 <= mo <= 12:
        return None
    nxt = date(y + (mo == 12), (mo % 12) + 1, 1)
    return date.fromordinal(nxt.toordinal() - 1).isoformat()


def _num(v):
    """엑셀은 같은 열에서도 숫자와 문자열이 섞여 온다."""
    if v is None or v == "":
        return None
    try:
        return round(float(str(v).strip()), 3)
    except ValueError:
        return None


def parse_xlsx(blob: bytes, sheet: str = "Data") -> list[dict]:
    """엑셀 바이트 → 시리즈 목록. 수집과 분리해 저장된 파일로 테스트할 수 있게 한다."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(blob), data_only=True, read_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    hdr = [str(h).strip() if h is not None else "" for h in next(rows)]
    # 헤더 위치를 이름으로 찾는다 (열 순서가 바뀌어도 견디게)
    pos = {h: i for i, h in enumerate(hdr)}
    missing = [h for h in XL_COLS if h not in pos]
    if missing:
        raise FrbsfError(f"엑셀에 없는 열: {missing[:3]} ... (열 이름이 바뀌었는지 확인)")
    if "time_month" not in pos:
        raise FrbsfError("엑셀에 time_month 열이 없습니다.")
    acc = {name: [] for name in XL_COLS.values()}
    for r in rows:
        d = _to_date(r[pos["time_month"]])
        if d is None:
            continue
        for col, name in XL_COLS.items():
            v = _num(r[pos[col]])
            if v is not None:
                acc[name].append({"d": d, "v": v})
    return [{"name": n, "data": sorted(acc[n], key=lambda p: p["d"])}
            for n in XL_COLS.values() if acc[n]]


def parse_csv(text: str, prefix: str = "") -> list[dict]:
    """차트 CSV 본문 → 시리즈 목록 (예비 경로)."""
    acc, order = {}, []
    for r in csv.DictReader(io.StringIO(text.lstrip("﻿"))):
        d = _to_date(r.get("time_month"))
        if d is None:
            continue
        for col, ko in CSV_COLS.items():
            v = _num(r.get(col))
            if v is None:
                continue
            name = f"{prefix} · {ko}" if prefix else ko
            if name not in acc:
                acc[name] = []
                order.append(name)
            acc[name].append({"d": d, "v": v})
    return [{"name": n, "data": sorted(acc[n], key=lambda p: p["d"])} for n in order]


def fetch(indicator: dict) -> list[dict]:
    """indicators.yaml 지표 하나 → 시리즈 목록 (다른 수집기와 동일 형식)."""
    p = indicator.get("params") or {}
    iid = indicator.get("id", "?")

    # ① 엑셀 (1969년~). 내려받기 실패하면 로컬 사본을 쓴다.
    name = p.get("xlsx", "supply-demand-pce-inflation.xlsx")
    url = name if name.startswith("http") else BASE + name
    blob = None
    try:
        r = requests.get(url, timeout=90, headers=UA)
        r.raise_for_status()
        blob = r.content
    except Exception as e:
        print(f"  [frbsf {iid}] 엑셀 내려받기 실패: {str(e)[:140]}")
        local = ROOT / (p.get("local") or "data/supply-demand-pce-inflation.xlsx")
        if local.exists():
            print(f"  [frbsf {iid}] 로컬 사본 사용: {local.name}")
            blob = local.read_bytes()
    if blob:
        try:
            out = parse_xlsx(blob, p.get("sheet", "Data"))
            n = sum(len(s["data"]) for s in out)
            span = min(s["data"][0]["d"] for s in out) + " ~ " + max(s["data"][-1]["d"] for s in out)
            print(f"  [frbsf {iid}] 엑셀 → 시리즈 {len(out)}개, 관측치 {n}개 ({span})")
            return out
        except Exception as e:
            print(f"  [frbsf {iid}] 엑셀 해석 실패: {str(e)[:140]}")

    # ② 예비 — 차트 CSV (최근 5년만)
    out = []
    for f in p.get("files") or []:
        u = f["csv"] if f["csv"].startswith("http") else BASE + f["csv"]
        try:
            r = requests.get(u, timeout=60, headers=UA)
            r.raise_for_status()
            got = parse_csv(r.text, f.get("prefix") or "")
        except Exception as e:
            print(f"  [frbsf {iid}] {f['csv']} 실패(건너뜀): {str(e)[:140]}")
            continue
        print(f"  [frbsf {iid}] {f['csv']} → 시리즈 {len(got)}개 (CSV 예비 경로)")
        out.extend(got)
    if not out:
        raise FrbsfError(f"FRBSF 수집 실패 ({iid}) — 엑셀·CSV 모두 받지 못했습니다.")
    return out
